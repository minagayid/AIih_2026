#!/usr/bin/env python3
"""Train an exploratory image-feature regressor from model-generated ratings."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageOps


FEATURE_NAMES = [
    "log_width",
    "log_height",
    "aspect_ratio",
    "gray_mean",
    "gray_std",
    "gray_p05",
    "gray_p25",
    "gray_p50",
    "gray_p75",
    "gray_p95",
    "gray_entropy_normalized",
    "laplacian_variance",
    "mean_abs_gradient_x",
    "mean_abs_gradient_y",
    "gradient_rms",
    "near_black_fraction",
    "near_white_fraction",
    *[f"gray_histogram_{i:02d}" for i in range(32)],
]
ALPHAS = [0.1, 1.0, 10.0, 100.0]


def image_features(image_bytes: bytes) -> np.ndarray:
    image = Image.open(BytesIO(image_bytes))
    width, height = image.size
    gray = np.asarray(
        ImageOps.grayscale(image).resize((96, 96), Image.Resampling.LANCZOS),
        dtype=np.float32,
    ) / 255.0
    histogram, _ = np.histogram(gray, bins=32, range=(0.0, 1.0))
    histogram = histogram.astype(np.float64) / histogram.sum()
    nonzero = histogram[histogram > 0]
    entropy = float(-(nonzero * np.log2(nonzero)).sum() / 5.0)

    gradient_x = gray[:, 1:] - gray[:, :-1]
    gradient_y = gray[1:, :] - gray[:-1, :]
    laplacian = (
        gray[1:-1, :-2]
        + gray[1:-1, 2:]
        + gray[:-2, 1:-1]
        + gray[2:, 1:-1]
        - 4.0 * gray[1:-1, 1:-1]
    )
    gradient_rms = np.sqrt(
        (gradient_x[:-1, :] ** 2 + gradient_y[:, :-1] ** 2) / 2.0
    )
    summary = np.array(
        [
            np.log1p(width),
            np.log1p(height),
            width / max(height, 1),
            gray.mean(),
            gray.std(),
            *np.quantile(gray, [0.05, 0.25, 0.50, 0.75, 0.95]),
            entropy,
            np.var(laplacian),
            np.mean(np.abs(gradient_x)),
            np.mean(np.abs(gradient_y)),
            np.sqrt(np.mean(gradient_rms**2)),
            np.mean(gray < 0.05),
            np.mean(gray > 0.95),
        ],
        dtype=np.float64,
    )
    return np.concatenate([summary, histogram])


def rankdata(values: np.ndarray) -> np.ndarray:
    order = np.argsort(values, kind="mergesort")
    ranks = np.empty(len(values), dtype=np.float64)
    start = 0
    while start < len(values):
        end = start + 1
        while end < len(values) and values[order[end]] == values[order[start]]:
            end += 1
        ranks[order[start:end]] = (start + 1 + end) / 2.0
        start = end
    return ranks


def spearman(left: np.ndarray, right: np.ndarray) -> float:
    return float(np.corrcoef(rankdata(left), rankdata(right))[0, 1])


def grouped_splits(
    indices: np.ndarray,
    groups: np.ndarray,
    fold_count: int,
    seed: int,
) -> list[tuple[np.ndarray, np.ndarray]]:
    unique_groups = np.unique(groups[indices])
    if len(unique_groups) < fold_count:
        raise ValueError("There are fewer unique images than requested folds.")
    shuffled = unique_groups.copy()
    np.random.default_rng(seed).shuffle(shuffled)
    folds = np.array_split(shuffled, fold_count)
    splits = []
    for test_groups in folds:
        test_mask = np.isin(groups[indices], test_groups)
        splits.append((indices[~test_mask], indices[test_mask]))
    return splits


def fit_ridge(
    features: np.ndarray,
    target: np.ndarray,
    alpha: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    mean = features.mean(axis=0)
    scale = features.std(axis=0)
    scale[scale < 1e-8] = 1.0
    standardized = (features - mean) / scale
    target_mean = float(target.mean())
    coefficients = np.linalg.solve(
        standardized.T @ standardized + alpha * np.eye(standardized.shape[1]),
        standardized.T @ (target - target_mean),
    )
    return mean, scale, coefficients, target_mean


def predict_ridge(
    features: np.ndarray,
    model: tuple[np.ndarray, np.ndarray, np.ndarray, float],
) -> np.ndarray:
    mean, scale, coefficients, target_mean = model
    return ((features - mean) / scale) @ coefficients + target_mean


def select_alpha(
    features: np.ndarray,
    target: np.ndarray,
    groups: np.ndarray,
    indices: np.ndarray,
    fold_count: int,
    seed: int,
) -> tuple[float, dict[str, float]]:
    splits = grouped_splits(indices, groups, fold_count, seed)
    losses = {}
    for alpha in ALPHAS:
        total_error = 0.0
        total_count = 0
        for train_indices, validation_indices in splits:
            model = fit_ridge(features[train_indices], target[train_indices], alpha)
            prediction = predict_ridge(features[validation_indices], model)
            total_error += float(
                np.abs(target[validation_indices] - prediction).sum()
            )
            total_count += len(validation_indices)
        losses[str(alpha)] = total_error / total_count
    best_alpha = min(ALPHAS, key=lambda alpha: losses[str(alpha)])
    return best_alpha, losses


def load_dataset(
    input_path: Path, sheet_name: str
) -> tuple[
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    list[str],
    int,
    dict[str, int],
]:
    workbook = load_workbook(input_path, data_only=True, read_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet {sheet_name!r} was not found. Available sheets: "
            + ", ".join(workbook.sheetnames)
        )
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The worksheet is empty.")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    score_columns = list(range(2, len(headers)))
    if not score_columns:
        raise ValueError(
            "Expected an Index column, an Image column, and model score columns."
        )
    score_names = [headers[index] or f"model_{index - 1}" for index in score_columns]

    grouped_scores: dict[str, list[np.ndarray]] = defaultdict(list)
    grouped_features: dict[str, np.ndarray] = {}
    dimensions: Counter[str] = Counter()
    image_rows = 0
    for embedded_image in worksheet._images:
        excel_row = embedded_image.anchor._from.row + 1
        if excel_row < 2 or excel_row > len(rows):
            continue
        row = rows[excel_row - 1]
        try:
            ratings = np.array(
                [float(row[index]) for index in score_columns], dtype=np.float64
            )
        except (IndexError, TypeError, ValueError) as error:
            raise ValueError(
                f"An image row near worksheet row {excel_row} has missing or "
                "non-numeric model ratings."
            ) from error
        image_bytes = embedded_image._data()
        digest = sha256(image_bytes).hexdigest()
        with Image.open(BytesIO(image_bytes)) as image:
            dimensions[f"{image.width}x{image.height}"] += 1
        grouped_scores[digest].append(ratings)
        grouped_features.setdefault(digest, image_features(image_bytes))
        image_rows += 1

    if not grouped_scores:
        raise ValueError("No embedded images paired with model ratings were found.")

    hashes = np.array(sorted(grouped_scores))
    rating_matrix = np.vstack(
        [np.mean(grouped_scores[digest], axis=0) for digest in hashes]
    )
    features = np.vstack([grouped_features[digest] for digest in hashes])
    target = rating_matrix.mean(axis=1)
    return (
        features,
        target,
        rating_matrix,
        hashes,
        score_names,
        image_rows,
        dict(dimensions),
    )


def evaluate(
    features: np.ndarray,
    target: np.ndarray,
    groups: np.ndarray,
    seed: int,
    folds: int,
) -> tuple[dict[str, object], np.ndarray]:
    all_indices = np.arange(len(target))
    outer_splits = grouped_splits(all_indices, groups, folds, seed)
    predictions = np.zeros(len(target), dtype=np.float64)
    baseline = np.zeros(len(target), dtype=np.float64)
    selected_alphas = []
    for fold_index, (train_indices, test_indices) in enumerate(outer_splits):
        alpha, _ = select_alpha(
            features,
            target,
            groups,
            train_indices,
            fold_count=min(4, len(np.unique(groups[train_indices]))),
            seed=seed + fold_index + 1,
        )
        model = fit_ridge(features[train_indices], target[train_indices], alpha)
        predictions[test_indices] = predict_ridge(features[test_indices], model)
        baseline[test_indices] = float(target[train_indices].mean())
        selected_alphas.append(alpha)

    residuals = target - predictions
    baseline_residuals = target - baseline
    summary: dict[str, object] = {
        "method": "Nested grouped cross-validation by exact image SHA-256",
        "outer_folds": folds,
        "alpha_candidates": ALPHAS,
        "selected_alpha_by_outer_fold": selected_alphas,
        "model_mae": float(np.mean(np.abs(residuals))),
        "mean_baseline_mae": float(np.mean(np.abs(baseline_residuals))),
        "model_rmse": float(np.sqrt(np.mean(residuals**2))),
        "mean_baseline_rmse": float(np.sqrt(np.mean(baseline_residuals**2))),
        "model_spearman_vs_weak_target": spearman(target, predictions),
        "r2_vs_weak_target": float(
            1.0 - np.sum(residuals**2) / np.sum((target - target.mean()) ** 2)
        ),
    }
    return summary, predictions


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Local .xlsx file")
    parser.add_argument("--sheet", default="Images", help="Worksheet name")
    parser.add_argument("--seed", type=int, default=20260927)
    parser.add_argument("--folds", type=int, default=5)
    parser.add_argument(
        "--results", type=Path, default=Path("ml/results.json")
    )
    parser.add_argument(
        "--model", type=Path, default=Path("work/quality_baseline_model.json")
    )
    args = parser.parse_args()

    (
        features,
        target,
        ratings,
        groups,
        rating_names,
        image_rows,
        image_dimensions,
    ) = load_dataset(args.input, args.sheet)
    rating_correlations = [
        spearman(ratings[:, left], ratings[:, right])
        for left in range(ratings.shape[1])
        for right in range(left + 1, ratings.shape[1])
    ]
    cv_summary, _ = evaluate(features, target, groups, args.seed, args.folds)
    final_alpha, final_alpha_losses = select_alpha(
        features,
        target,
        groups,
        np.arange(len(target)),
        fold_count=args.folds,
        seed=args.seed + 100,
    )
    final_model = fit_ridge(features, target, final_alpha)

    summary = {
        "dataset": {
            "worksheet": args.sheet,
            "scored_image_rows": image_rows,
            "exact_unique_images": len(target),
            "embedded_image_dimensions": image_dimensions,
            "model_rating_columns": len(rating_names),
            "rating_scale_documented_in_source": False,
            "expert_reference_scores_present": False,
        },
        "target": "Arithmetic mean of the seven supplied model ratings per exact-unique image; weak label only.",
        "model": {
            "type": "Ridge regression",
            "input": "49 handcrafted grayscale, histogram, and image-quality features from a 96x96 resize",
            "feature_count": len(FEATURE_NAMES),
            "seed": args.seed,
            "selected_alpha_on_full_data": final_alpha,
            "full_data_selection_mae_by_alpha": final_alpha_losses,
        },
        "model_rating_agreement": {
            "mean_pairwise_spearman": float(np.mean(rating_correlations)),
            "median_pairwise_spearman": float(np.median(rating_correlations)),
            "pair_count": len(rating_correlations),
        },
        "cross_validation": cv_summary,
        "limitations": [
            "No expert reference scores were present; validation measures reproduction of model-generated ratings only.",
            "The workbook did not document a rating scale or scoring rubric.",
            "The workbook schema had no patient identifiers, so patient-level separation could not be verified.",
            "Images were embedded at thumbnail dimensions, not clinical acquisition resolution.",
            "Exact byte-identical images were aggregated before training; near-duplicates and same-patient images were not independently verified.",
            "This exploratory model is not validated for diagnosis, retake decisions, or clinical use.",
        ],
    }
    model_artifact = {
        "model_type": "ridge_regression",
        "target": summary["target"],
        "feature_names": FEATURE_NAMES,
        "feature_mean": final_model[0].tolist(),
        "feature_scale": final_model[1].tolist(),
        "standardized_coefficients": final_model[2].tolist(),
        "target_mean": final_model[3],
        "alpha": final_alpha,
        "training_samples": len(target),
        "contains_patient_images_or_row_data": False,
        "for_clinical_use": False,
    }

    args.results.parent.mkdir(parents=True, exist_ok=True)
    args.model.parent.mkdir(parents=True, exist_ok=True)
    args.results.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    args.model.write_text(
        json.dumps(model_artifact, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
